#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Su Ayak İzi Sektör Kataloğu
=============================
Sektörlere ve süreçlere göre tipik su yoğunluğu, baskın su ayak izi türü
(mavi/yeşil/gri su) ve başlıca su risklerinin; ISO 14046, AWS Standard,
CDP Su Güvenliği ve GRI 303 gibi başlıca çerçevelere göre düzenlenmiş,
aranabilir ve filtrelenebilir kataloğu.

DOSYA YAPISI
------------
    suayakizi.py                     Uygulama (bu dosya)
    su_ayak_izi_sektorleri.json      Sektör/süreç veri seti (koddan ayrı)
    static/icons/*.png               PWA ikonları

ÇALIŞTIRMA (geliştirme / Termux dahil, herhangi bir Python 3 ortamında)
------------------------------------------------------------------------
    pip install -r requirements.txt      (yalnızca ilk seferde)
    python suayakizi.py

Sonra telefonun (veya bilgisayarın) tarayıcısından şu adresi aç:
    http://127.0.0.1:8081

PROD ORTAMI
-----------
    pip install -r requirements.txt gunicorn
    gunicorn -w 2 -b 0.0.0.0:8081 suayakizi:app

Docker ile:
    docker build -t su-ayak-izi-katalogu .
    docker run -p 8081:8081 su-ayak-izi-katalogu

Testler:
    pip install -r requirements-dev.txt
    pytest

Notlar
------
* Veri seti kod içine gömülü DEĞİLDİR; su_ayak_izi_sektorleri.json
  dosyasından okunur ve başlangıçta doğrulanır.
* Arama ve sektör/tür filtresi tamamen istemci tarafında çalışır.
* Bu uygulama, diğer kardeş projelerinden (CBAM Kataloğu, Karbon
  Kredisi Kataloğu, CSDDD Kontrolcüsü) farklı olarak varsayılan olarak
  AÇIK, su temalı bir arayüzle açılır; koyu "derin deniz" teması
  isteğe bağlı olarak seçilebilir.
* Filtrelenmiş sonuçlar, kurumsal biçimli renk hiyerarşili bir Excel
  (.xlsx) raporu olarak /export/xlsx uç noktasından indirilebilir.
* Bu katalog genel bilgilendirme amaçlıdır; sektörel su yoğunluğu
  düzeyleri göreli/nitel kategoriler olarak sunulur, kesin ölçüm
  değeri yerine geçmez.
"""

import io
import json
import logging
import math
import os
import sys
import webbrowser
from datetime import datetime
from threading import Timer

from flask import Flask, Response, jsonify, request, send_from_directory
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

# =============================================================================
# 0. TEMEL AYARLAR VE LOGLAMA
# =============================================================================
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_FILE = os.path.join(BASE_DIR, "su_ayak_izi_sektorleri.json")
ICONS_DIR = os.path.join(BASE_DIR, "static", "icons")
PORT = int(os.environ.get("SUAYAK_PORT", 8081))
HOST = os.environ.get("SUAYAK_HOST", "127.0.0.1")

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)
log = logging.getLogger("suayakizi")

REQUIRED_FIELDS = (
    "Sektör", "Süreç", "Su Yoğunluğu", "Su Ayak İzi Türü",
    "Başlıca Su Riski", "İlgili Standartlar", "Azaltım Uygulamaları", "Açıklama",
)
ALLOWED_INTENSITY = ("Düşük", "Orta", "Yüksek", "Çok Yüksek")
ALLOWED_TYPES = ("Mavi Su", "Yeşil Su", "Gri Su", "Karma")


# =============================================================================
# 1. VERİ YÜKLEME VE DOĞRULAMA
# =============================================================================
def _unique_ordered(seq):
    """Bir dizideki benzersiz değerleri, ilk görüldükleri sırayla döndürür.

    Sektör adları "1. ...", "2. ...", ... "10. ..." biçiminde numaralı
    olduğundan alfabetik sıralama yanlış sonuç verir; bu yüzden veri
    dosyasındaki doğal sırayı korur.
    """
    seen = []
    for item in seq:
        if item not in seen:
            seen.append(item)
    return seen


def load_data(path):
    """su_ayak_izi_sektorleri.json dosyasını okur, doğrular ve döndürür."""
    if not os.path.isfile(path):
        log.error("Veri dosyası bulunamadı: %s", path)
        log.error("su_ayak_izi_sektorleri.json dosyasının suayakizi.py ile aynı klasörde olduğundan emin olun.")
        sys.exit(1)

    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
    except json.JSONDecodeError as exc:
        log.error("Veri dosyası geçerli bir JSON değil (%s): %s", path, exc)
        sys.exit(1)

    if not isinstance(data, list) or not data:
        log.error("Veri dosyası boş veya beklenen liste formatında değil.")
        sys.exit(1)

    errors = []
    for i, row in enumerate(data):
        if not isinstance(row, dict):
            errors.append(f"Kayıt #{i}: sözlük (obje) formatında değil.")
            continue
        missing = [field for field in REQUIRED_FIELDS if not row.get(field)]
        if missing:
            errors.append(f"Kayıt #{i} ({row.get('Süreç', '?')}): eksik alan(lar): {', '.join(missing)}")
            continue
        if row["Su Yoğunluğu"] not in ALLOWED_INTENSITY:
            errors.append(
                f"Kayıt #{i} ({row['Süreç']}): geçersiz Su Yoğunluğu '{row['Su Yoğunluğu']}' "
                f"(izin verilenler: {', '.join(ALLOWED_INTENSITY)})"
            )
        if row["Su Ayak İzi Türü"] not in ALLOWED_TYPES:
            errors.append(
                f"Kayıt #{i} ({row['Süreç']}): geçersiz Su Ayak İzi Türü '{row['Su Ayak İzi Türü']}' "
                f"(izin verilenler: {', '.join(ALLOWED_TYPES)})"
            )

    if errors:
        log.error("Veri doğrulaması başarısız oldu (%d hata):", len(errors))
        for e in errors[:20]:
            log.error("  - %s", e)
        sys.exit(1)

    log.info("Veri yüklendi: %d kayıt, kaynak: %s", len(data), path)
    return data


DATA = load_data(DATA_FILE)
SECTORS = _unique_ordered(row["Sektör"] for row in DATA)
WATER_TYPES = [t for t in ALLOWED_TYPES if t in set(row["Su Ayak İzi Türü"] for row in DATA)]


# =============================================================================
# 1b. XLSX (EXCEL) DIŞA AKTARIMI
# =============================================================================
_XLSX_SECTOR_COLORS = {
    "1. TEKSTİL VE HAZIR GİYİM": "7E22CE",
    "2. GIDA VE İÇECEK": "A16207",
    "3. TARIM VE HAYVANCILIK": "4D7C0F",
    "4. MADENCİLİK VE METALLER": "57534E",
    "5. ENERJİ ÜRETİMİ": "C2410C",
    "6. KAĞIT VE SELÜLOZ": "15803D",
    "7. KİMYA VE PETROKİMYA": "B91C1C",
    "8. ELEKTRONİK VE YARI İLETKEN": "4338CA",
    "9. DERİ İŞLEME (TABAKLAMA)": "92400E",
    "10. VERİ MERKEZLERİ (SOĞUTMA)": "0369A1",
}
_XLSX_INTENSITY_STYLES = {
    "Düşük": {"fill": "D1FAE5", "font": "065F46"},
    "Orta": {"fill": "FEF9C3", "font": "854D0E"},
    "Yüksek": {"fill": "FED7AA", "font": "9A3412"},
    "Çok Yüksek": {"fill": "FEE2E2", "font": "991B1B"},
}
_XLSX_TYPE_COLORS = {
    "Mavi Su": "1D4ED8", "Yeşil Su": "15803D", "Gri Su": "57534E", "Karma": "7E22CE",
}

_XLSX_NAVY_DARK = "0F172A"
_XLSX_NAVY = "1E293B"
_XLSX_ACCENT = "0E7490"
_XLSX_BORDER = "E2E8F0"
_XLSX_ZEBRA = "F2FAFC"
_XLSX_TEXT_DARK = "111827"
_XLSX_TEXT_MUTE = "64748B"
_XLSX_WHITE = "FFFFFF"

_XLSX_COLUMNS = [
    "Sektör", "Süreç", "Su Yoğunluğu", "Su Ayak İzi Türü",
    "Başlıca Su Riski", "İlgili Standartlar", "Azaltım Uygulamaları", "Açıklama",
]
_XLSX_COL_WIDTHS = {"A": 30, "B": 32, "C": 13, "D": 13, "E": 38, "F": 26, "G": 32, "H": 42}
_XLSX_HEADER_ROW = 5
_XLSX_CHARS_PER_LINE = {"A": 28, "B": 30, "E": 36, "F": 24, "G": 30, "H": 40}


def _xlsx_border(color=_XLSX_BORDER):
    side = Side(style="thin", color=color)
    return Border(left=side, right=side, top=side, bottom=side)


def _build_report_xlsx(rows, sector_filter, type_filter, query):
    """Filtrelenmiş kayıtlardan kurumsal görünümlü Excel raporu üretir."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Su Ayak İzi Kataloğu"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = _XLSX_ACCENT

    for col, width in _XLSX_COL_WIDTHS.items():
        ws.column_dimensions[col].width = width

    last_col = get_column_letter(len(_XLSX_COLUMNS))

    ws.merge_cells(f"A1:{last_col}1")
    cell = ws["A1"]
    cell.value = "Su Ayak İzi Sektör Kataloğu"
    cell.font = Font(name="Calibri", size=17, bold=True, color=_XLSX_WHITE)
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for i in range(len(_XLSX_COLUMNS)):
        ws[f"{get_column_letter(i + 1)}1"].fill = PatternFill("solid", fgColor=_XLSX_NAVY_DARK)
    ws.row_dimensions[1].height = 34

    ws.merge_cells(f"A2:{last_col}2")
    cell = ws["A2"]
    cell.value = "Sektöre ve sürece göre su yoğunluğu, su ayak izi türü ve başlıca su riski; ISO 14046, AWS Standard, CDP Su Güvenliği ve GRI 303 çerçevelerine göre düzenlenmiştir"
    cell.font = Font(name="Calibri", size=10, italic=True, color="CBD5E1")
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for i in range(len(_XLSX_COLUMNS)):
        ws[f"{get_column_letter(i + 1)}2"].fill = PatternFill("solid", fgColor=_XLSX_NAVY_DARK)
    ws.row_dimensions[2].height = 20

    filt_bits = [f"Sektör: {sector_filter or 'Tümü'}", f"Tür: {type_filter or 'Tümü'}"]
    if query:
        filt_bits.append(f'Arama: "{query}"')
    meta_text = (
        f"Oluşturma: {datetime.now().strftime('%d.%m.%Y %H:%M')}   \u2022   "
        f"Kayıt: {len(rows)} / {len(DATA)}   \u2022   " + "   \u00b7   ".join(filt_bits)
    )
    ws.merge_cells(f"A3:{last_col}3")
    cell = ws["A3"]
    cell.value = meta_text
    cell.font = Font(name="Calibri", size=9, color="94A3B8")
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for i in range(len(_XLSX_COLUMNS)):
        ws[f"{get_column_letter(i + 1)}3"].fill = PatternFill("solid", fgColor=_XLSX_NAVY)
    ws.row_dimensions[3].height = 18

    ws.row_dimensions[4].height = 8

    for i, header in enumerate(_XLSX_COLUMNS):
        col = get_column_letter(i + 1)
        cell = ws[f"{col}{_XLSX_HEADER_ROW}"]
        cell.value = header
        cell.font = Font(name="Calibri", size=11, bold=True, color=_XLSX_WHITE)
        cell.fill = PatternFill("solid", fgColor=_XLSX_ACCENT)
        cell.alignment = Alignment(
            horizontal="center" if col in ("C", "D") else "left",
            vertical="center", indent=0 if col in ("C", "D") else 1,
        )
        cell.border = _xlsx_border(_XLSX_ACCENT)
    ws.row_dimensions[_XLSX_HEADER_ROW].height = 24

    r = _XLSX_HEADER_ROW + 1
    for idx, row in enumerate(rows):
        row_fill = PatternFill("solid", fgColor=_XLSX_ZEBRA if idx % 2 == 1 else _XLSX_WHITE)

        sektor = row["Sektör"]
        surec = row["Süreç"]
        yogunluk = row["Su Yoğunluğu"]
        tur = row["Su Ayak İzi Türü"]
        risk = row["Başlıca Su Riski"]
        standartlar = row["İlgili Standartlar"]
        azaltim = row["Azaltım Uygulamaları"]
        aciklama = row["Açıklama"]

        sc = _XLSX_SECTOR_COLORS.get(sektor, _XLSX_TEXT_MUTE)
        sektor_kisa = sektor.split(". ", 1)[-1] if ". " in sektor else sektor
        cell = ws.cell(row=r, column=1, value=CellRichText(
            TextBlock(InlineFont(rFont="Calibri", b=True, color=sc), "\u25a0 "),
            TextBlock(InlineFont(rFont="Calibri", color=_XLSX_TEXT_DARK), sektor_kisa),
        ))
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
        cell.fill = row_fill
        cell.border = _xlsx_border()

        cell = ws.cell(row=r, column=2, value=surec)
        cell.font = Font(name="Calibri", size=10.5, bold=True, color=_XLSX_TEXT_DARK)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
        cell.fill = row_fill
        cell.border = _xlsx_border()

        istyle = _XLSX_INTENSITY_STYLES.get(yogunluk, {"fill": "F1F5F9", "font": _XLSX_TEXT_MUTE})
        cell = ws.cell(row=r, column=3, value=yogunluk)
        cell.font = Font(name="Calibri", size=9.5, bold=True, color=istyle["font"])
        cell.fill = PatternFill("solid", fgColor=istyle["fill"])
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _xlsx_border()

        tc = _XLSX_TYPE_COLORS.get(tur, _XLSX_TEXT_MUTE)
        cell = ws.cell(row=r, column=4, value=tur)
        cell.font = Font(name="Calibri", size=9.5, bold=True, color=tc)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = row_fill
        cell.border = _xlsx_border()

        cell = ws.cell(row=r, column=5, value=risk)
        cell.font = Font(name="Calibri", size=9.5, color=_XLSX_TEXT_DARK)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
        cell.fill = row_fill
        cell.border = _xlsx_border()

        cell = ws.cell(row=r, column=6, value=standartlar)
        cell.font = Font(name="Calibri", size=9.5, color=_XLSX_TEXT_DARK)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
        cell.fill = row_fill
        cell.border = _xlsx_border()

        cell = ws.cell(row=r, column=7, value=azaltim)
        cell.font = Font(name="Calibri", size=9.5, color=_XLSX_TEXT_DARK)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
        cell.fill = row_fill
        cell.border = _xlsx_border()

        cell = ws.cell(row=r, column=8, value=aciklama)
        cell.font = Font(name="Calibri", size=9.5, color=_XLSX_TEXT_DARK)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
        cell.fill = row_fill
        cell.border = _xlsx_border()

        candidates = [
            math.ceil(len(sektor_kisa) / _XLSX_CHARS_PER_LINE["A"]),
            math.ceil(len(surec) / _XLSX_CHARS_PER_LINE["B"]),
            math.ceil(len(risk) / _XLSX_CHARS_PER_LINE["E"]),
            math.ceil(len(standartlar) / _XLSX_CHARS_PER_LINE["F"]),
            math.ceil(len(azaltim) / _XLSX_CHARS_PER_LINE["G"]),
            math.ceil(len(aciklama) / _XLSX_CHARS_PER_LINE["H"]),
        ]
        lines = max(1, *candidates)
        ws.row_dimensions[r].height = max(18, lines * 14 + 6)
        r += 1

    last_row = r - 1
    if rows:
        ws.auto_filter.ref = f"A{_XLSX_HEADER_ROW}:{last_col}{last_row}"
    ws.freeze_panes = f"A{_XLSX_HEADER_ROW + 1}"

    footer_row = last_row + 2
    ws.merge_cells(f"A{footer_row}:{last_col}{footer_row}")
    cell = ws.cell(row=footer_row, column=1)
    cell.value = (
        "Kaynak: Bu katalog, ISO 14046 (Su Ayak İzi), AWS Standard (Alliance for Water Stewardship), "
        "CDP Su Güvenliği bildirim çerçevesi ve GRI 303 (Su ve Atık Su) gibi başlıca çerçevelerde yaygın "
        "kullanılan kavramların genel bir özetidir. Su yoğunluğu düzeyleri göreli/nitel kategorilerdir, kesin "
        "ölçüm değeri yerine geçmez ve konum, teknoloji ve ölçeğe göre önemli ölçüde değişebilir. Bu içerik "
        "yatırım, hukuki veya teknik danışmanlık teşkil etmez; saha bazlı değerlendirme için yetkin bir "
        "danışmana veya ilgili standart kuruluşunun resmi dokümanlarına başvurulmalıdır."
    )
    cell.font = Font(name="Calibri", size=8.5, italic=True, color=_XLSX_TEXT_MUTE)
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    ws.row_dimensions[footer_row].height = 48

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{_XLSX_HEADER_ROW}:{_XLSX_HEADER_ROW}"
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5, header=0.2, footer=0.2)

    wb.properties.title = "Su Ayak İzi Sektör Kataloğu"
    wb.properties.creator = "Su Ayak İzi Katalog Uygulaması"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# =============================================================================
# 2. FLASK UYGULAMASI
# =============================================================================
app = Flask(__name__)

_PAGE_HTML = r"""<!doctype html>
<html lang="tr" data-theme="light">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">
<title>Su Ayak İzi Sektör Kataloğu</title>
<meta name="description" content="Sektörlere göre su yoğunluğu, su ayak izi türü ve su riski kataloğu: ISO 14046, AWS Standard, CDP Su Güvenliği ve GRI 303 çerçevelerine göre düzenlenmiştir">
<meta name="theme-color" content="#eaf7fb">
<meta name="mobile-web-app-capable" content="yes">
<meta name="apple-mobile-web-app-capable" content="yes">
<meta name="apple-mobile-web-app-status-bar-style" content="default">
<meta name="apple-mobile-web-app-title" content="Su Ayak İzi">
<link rel="manifest" href="/manifest.json">
<link rel="apple-touch-icon" href="/icons/icon-192.png">
<link rel="icon" href="/icons/icon-192.png">
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');

:root{
  --bg:#eaf7fb;
  --surface:#ffffff;
  --surface-2:#eef8fb;
  --border:#cfe8f0;
  --border-soft:#e3f3f8;
  --text:#0b2530;
  --text-dim:#35606b;
  --text-mute:#7ea3ac;
  --accent:#0891b2;
  --accent-strong:#06b6d4;
  --accent-soft:rgba(8,145,178,.12);
  --radius-lg:16px;
  --radius-md:10px;
  --radius-sm:7px;
}
[data-theme="dark"]{
  --bg:#061a20;
  --surface:#0d2530;
  --surface-2:#112e3a;
  --border:#17414f;
  --border-soft:#123340;
  --text:#e6f6fb;
  --text-dim:#9fc7d1;
  --text-mute:#5c8894;
  --accent:#22d3ee;
  --accent-strong:#67e8f9;
  --accent-soft:rgba(34,211,238,.16);
}
[data-theme="dark"] body{
  background:
    radial-gradient(1100px 480px at 50% -140px, #0e3644 0%, transparent 62%),
    var(--bg);
}
[data-theme="dark"] .topbar{background:rgba(6,26,32,.88);}
*{box-sizing:border-box;-webkit-tap-highlight-color:transparent;}
html,body{margin:0;padding:0;}
body{
  background:
    radial-gradient(1200px 520px at 50% -160px, #d3f0f7 0%, transparent 60%),
    linear-gradient(180deg, #eaf7fb 0%, #e4f3f9 55%, #eaf7fb 100%);
  color:var(--text);
  font-family:'Inter',-apple-system,BlinkMacSystemFont,'Segoe UI',Roboto,sans-serif;
  -webkit-font-smoothing:antialiased;
  min-height:100vh;
  padding-bottom:44px;
}
h1,h2,h3,p{margin:0;}
button{font-family:inherit;}

.topbar{
  position:sticky;top:0;z-index:30;
  display:flex;align-items:center;gap:13px;
  padding:15px 16px;
  background:rgba(255,255,255,.82);
  backdrop-filter:blur(16px) saturate(140%);
  -webkit-backdrop-filter:blur(16px) saturate(140%);
  border-bottom:1px solid var(--border-soft);
}
.brand-badge{
  width:40px;height:40px;border-radius:11px;flex-shrink:0;
  display:flex;align-items:center;justify-content:center;
  font-size:19px;
  background:linear-gradient(150deg,#083047,#0891b2);
  border:1px solid var(--border);
}
.brand h1{font-size:15.5px;font-weight:700;letter-spacing:-.01em;line-height:1.3;}
.brand p{font-size:11px;color:var(--text-mute);margin-top:3px;line-height:1.5;}

.wave-divider{display:block;width:100%;height:20px;margin-top:-1px;}
.wave-divider path{fill:var(--surface-2);}

.toolbar{padding:14px 16px 8px;}
.search-box{
  position:relative;display:flex;align-items:center;gap:8px;
  background:var(--surface);border:1px solid var(--border);
  border-radius:var(--radius-md);padding:0 12px;
  transition:border-color .15s;
}
.search-box:focus-within{border-color:var(--accent);}
.search-box svg{flex-shrink:0;color:var(--text-mute);}
.search-box input{
  flex:1;min-width:0;background:transparent;border:0;outline:0;
  color:var(--text);font-size:14.5px;padding:11px 2px;font-family:inherit;
}
.search-box input::placeholder{color:var(--text-mute);}
.clear-btn{
  display:none;align-items:center;justify-content:center;
  width:22px;height:22px;border-radius:50%;
  background:var(--surface-2);border:0;color:var(--text-dim);cursor:pointer;flex-shrink:0;
}
.clear-btn.show{display:flex;}

.chip-row{
  display:flex;gap:7px;overflow-x:auto;margin-top:10px;padding:1px 1px 4px;
  -ms-overflow-style:none;scrollbar-width:none;
}
.chip-row::-webkit-scrollbar{display:none;}
.chip{
  flex-shrink:0;display:flex;align-items:center;gap:6px;
  padding:7px 12px;border-radius:999px;
  border:1px solid var(--border);background:var(--surface);
  color:var(--text-dim);font-size:12.5px;font-weight:500;
  cursor:pointer;white-space:nowrap;transition:all .15s;
}
.chip .dot{width:7px;height:7px;border-radius:50%;flex-shrink:0;}
.chip .count{color:var(--text-mute);font-variant-numeric:tabular-nums;font-size:11.5px;}
.chip.active{background:var(--accent-soft);border-color:var(--accent);color:var(--text);}

.topbar-actions{display:flex;gap:6px;margin-left:auto;flex-shrink:0;}
.icon-btn{
  width:34px;height:34px;border-radius:9px;flex-shrink:0;
  display:flex;align-items:center;justify-content:center;
  background:var(--surface);border:1px solid var(--border);
  color:var(--text-dim);font-size:13px;font-weight:700;cursor:pointer;
  transition:color .15s,border-color .15s;
}
.icon-btn:hover{color:var(--text);border-color:var(--accent);}

.filter-row{display:flex;gap:8px;align-items:center;margin-top:9px;}
.type-select{
  background:var(--surface);border:1px solid var(--border);border-radius:var(--radius-sm);
  color:var(--text-dim);font-size:12px;padding:7px 8px;font-family:inherit;cursor:pointer;
}
.export-btn{
  display:flex;align-items:center;gap:5px;
  background:var(--surface);border:1px solid var(--border);border-radius:var(--radius-sm);
  color:var(--text-dim);font-size:11.5px;font-weight:600;padding:7px 10px;cursor:pointer;
  white-space:nowrap;margin-left:auto;transition:color .15s,border-color .15s;
}
.export-btn:hover{color:var(--text);border-color:var(--accent);}

.stats-row{display:flex;gap:8px;padding:6px 16px 12px;flex-wrap:wrap;}
.stat-pill{
  display:flex;align-items:baseline;gap:6px;
  background:var(--surface);border:1px solid var(--border-soft);
  border-radius:var(--radius-sm);padding:7px 11px;
  font-size:11.5px;color:var(--text-mute);
}
.stat-pill b{color:var(--text);font-size:13.5px;font-weight:700;}

.results{padding:2px 16px 6px;display:flex;flex-direction:column;gap:9px;}
.card{
  background:var(--surface);
  border:1px solid var(--border-soft);
  border-left:3px solid var(--sc,var(--accent));
  border-radius:var(--radius-md);
  padding:13px 14px;
  box-shadow:0 1px 2px rgba(8,60,80,.04);
}
.card-top{display:flex;justify-content:space-between;align-items:center;gap:10px;margin-bottom:8px;}
.card-cat{
  display:flex;align-items:center;gap:6px;
  font-size:10px;color:var(--text-mute);text-transform:uppercase;letter-spacing:.05em;font-weight:700;
}
.card-cat .dot{width:7px;height:7px;border-radius:50%;flex-shrink:0;}
.card-title{font-size:13.5px;font-weight:700;color:var(--text);line-height:1.45;margin-bottom:7px;}
.card-desc{font-size:12px;color:var(--text-dim);line-height:1.65;margin-bottom:10px;}
.card-meta{display:flex;flex-wrap:wrap;gap:6px 8px;align-items:center;margin-bottom:6px;}
.badge{font-size:10.5px;font-weight:700;padding:3px 8px;border-radius:999px;letter-spacing:.02em;white-space:nowrap;}
.detail{font-size:11.5px;color:var(--text-mute);line-height:1.6;display:block;}
.detail b{color:var(--text-dim);font-weight:600;}
.detail + .detail{margin-top:3px;}

.empty{text-align:center;padding:60px 24px;color:var(--text-mute);}
.empty svg{opacity:.35;margin-bottom:14px;}
.empty p{font-size:13.5px;margin:0 0 16px;}
.empty button{
  background:var(--surface);border:1px solid var(--border);color:var(--text);
  padding:9px 18px;border-radius:8px;font-size:12.5px;cursor:pointer;
}

.notice{margin:16px 16px 0;border-radius:var(--radius-md);border:1px solid var(--border-soft);overflow:hidden;}
.notice.amber{border-color:rgba(217,119,6,.28);background:rgba(217,119,6,.05);}
.notice.cyan{border-color:rgba(8,145,178,.28);background:rgba(8,145,178,.06);}
.notice summary{
  cursor:pointer;padding:13px 15px;font-size:12.5px;font-weight:700;
  display:flex;align-items:center;gap:9px;list-style:none;
}
.notice summary::-webkit-details-marker{display:none;}
.notice.amber summary{color:#b45309;}
.notice.cyan summary{color:#0e7490;}
[data-theme="dark"] .notice.amber summary{color:#fbbf6a;}
[data-theme="dark"] .notice.cyan summary{color:#67e8f9;}
.notice summary .chev{margin-left:auto;font-size:10px;opacity:.65;transition:transform .15s;}
.notice[open] summary .chev{transform:rotate(180deg);}
.notice-body{padding:0 15px 15px;font-size:12px;line-height:1.75;color:var(--text-dim);}
.notice-body b{color:var(--text);}
.notice-body .line{display:block;margin-top:7px;}

.footer{
  margin:18px 16px 0;padding:14px 15px;
  border:1px solid var(--border-soft);border-radius:var(--radius-md);
  font-size:10.5px;line-height:1.75;color:var(--text-mute);
  display:flex;gap:10px;
}
.footer b{color:var(--text-dim);}
.footer svg{flex-shrink:0;margin-top:1px;color:var(--text-mute);}

@media(min-width:760px){
  .topbar{padding:20px 30px;}
  .toolbar{padding-left:30px;padding-right:30px;}
  .toolbar,.stats-row,.results,.notice,.footer{max-width:860px;margin-left:auto;margin-right:auto;}
  .stats-row{padding-left:30px;padding-right:30px;}
  .results{padding-left:30px;padding-right:30px;}
}
</style>
</head>
<body>

  <header class="topbar">
    <div class="brand-badge">💧</div>
    <div class="brand">
      <h1 id="brandTitle">Su Ayak İzi Sektör Kataloğu</h1>
      <p id="brandSubtitle">Sektöre göre su yoğunluğu, su ayak izi türü ve su riski; ISO 14046, AWS Standard, CDP Su Güvenliği ve GRI 303 çerçevelerine göre düzenlenmiştir</p>
    </div>
    <div class="topbar-actions">
      <button class="icon-btn" id="langToggle" type="button" aria-label="Dili değiştir / Switch language" title="TR / EN">TR</button>
      <button class="icon-btn" id="themeToggle" type="button" aria-label="Koyu / açık temayı değiştir" title="Tema">🌙</button>
    </div>
  </header>

  <div class="toolbar" role="search">
    <div class="search-box">
      <svg width="17" height="17" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><circle cx="11" cy="11" r="8"/><line x1="21" y1="21" x2="16.65" y2="16.65"/></svg>
      <input id="search" type="text" placeholder="Sektör, süreç veya standart ara" autocomplete="off" spellcheck="false" aria-label="Katalogda ara">
      <button class="clear-btn" id="clearSearch" aria-label="Aramayı temizle">
        <svg width="12" height="12" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.5" stroke-linecap="round"><line x1="18" y1="6" x2="6" y2="18"/><line x1="6" y1="6" x2="18" y2="18"/></svg>
      </button>
    </div>
    <div class="chip-row" id="sectorChips" role="group" aria-label="Sektöre göre filtrele"></div>
    <div class="filter-row">
      <select class="type-select" id="typeFilter" aria-label="Su ayak izi türüne göre filtrele">
        <option value="Tümü" id="typeAllOpt">Tüm Türler</option>
        <option value="Mavi Su">Mavi Su</option>
        <option value="Yeşil Su">Yeşil Su</option>
        <option value="Gri Su">Gri Su</option>
        <option value="Karma">Karma</option>
      </select>
      <button class="export-btn" id="exportExcel" type="button">⬇ Excel İndir</button>
    </div>
  </div>

  <div class="stats-row" id="statsRow" aria-live="polite"></div>

  <main class="results" id="results"></main>

  <section class="notice cyan">
    <details open>
      <summary>
        <svg width="16" height="16" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="M12 2.69l5.66 5.66a8 8 0 1 1-11.31 0z"/></svg>
        Su Ayak İzi Türleri Nedir?
        <span class="chev">▾</span>
      </summary>
      <div class="notice-body">
        <span class="line">• <b>Mavi Su:</b> Sulama, proses veya soğutma için yerüstü ve yeraltı kaynaklarından çekilip tüketilen (buharlaşan, ürüne karışan veya farklı bir havzaya aktarılan) tatlı su.</span>
        <span class="line">• <b>Yeşil Su:</b> Toprakta depolanıp bitkiler tarafından kullanılan yağmur suyu; özellikle yem ve tahıl üretiminde baskındır.</span>
        <span class="line">• <b>Gri Su:</b> Oluşan kirletici yükün, alıcı ortam kalite standartlarını aşmadan özümsenmesi için gereken varsayımsal tatlı su miktarı; yüksek kirletici deşarjı olan proseslerde belirleyicidir.</span>
        <span class="line">• Bu üçlü ayrım, Su Ayak İzi Ağı (Water Footprint Network) metodolojisine ve ISO 14046 standardına dayanır.</span>
      </div>
    </details>
  </section>

  <section class="notice amber">
    <details>
      <summary>
        <svg width="16" height="16" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><circle cx="12" cy="12" r="10"/><line x1="12" y1="8" x2="12" y2="12"/><line x1="12" y1="16" x2="12.01" y2="16"/></svg>
        Su Riski Değerlendirirken Nelere Dikkat Edilmeli?
        <span class="chev">▾</span>
      </summary>
      <div class="notice-body">
        <span class="line">• <b>Yoğunluk tek başına yeterli değildir:</b> Aynı su yoğunluğuna sahip bir tesis, bol sulu bir havzada düşük risk taşırken, su stresi yüksek bir bölgede ciddi risk taşıyabilir; konum bağlamı belirleyicidir.</span>
        <span class="line">• <b>Üç risk türü:</b> Fiziksel risk (su kıtlığı/kuraklık), düzenleyici risk (kullanım kısıtlaması, artan tarife) ve itibar riski (yerel toplulukla su paylaşımı anlaşmazlığı) birlikte değerlendirilmelidir.</span>
        <span class="line">• <b>Dolaylı (tedarik zinciri) su ayak izi:</b> Bir şirketin doğrudan su kullanımı genellikle toplam su ayak izinin küçük bir kısmıdır; büyük kısım hammadde tedarik zincirinde (özellikle tarımsal girdilerde) gizlidir.</span>
        <span class="line">• <b>Havza bazlı bakış:</b> Su, karbondan farklı olarak yerel bir kaynaktır; aynı litre sayısı farklı havzalarda çok farklı anlam taşır, bu nedenle konum bazlı risk haritalama araçları (ör. WRI Aqueduct) önerilir.</span>
      </div>
    </details>
  </section>

  <footer class="footer">
    <svg width="18" height="18" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><circle cx="12" cy="12" r="10"/><line x1="12" y1="16" x2="12" y2="12"/><line x1="12" y1="8" x2="12.01" y2="8"/></svg>
    <div>
      <b>Kaynak:</b> Bu katalog, ISO 14046 (Su Ayak İzi), AWS Standard, CDP Su Güvenliği bildirim çerçevesi ve GRI 303 gibi başlıca çerçevelerde yaygın kullanılan kavramların genel bir özetidir.<br><br>
      <b>Sorumluluk Reddi:</b> Su yoğunluğu düzeyleri göreli/nitel kategorilerdir, kesin ölçüm değeri yerine geçmez ve konum, teknoloji ve ölçeğe göre değişebilir. Bu içerik yatırım, hukuki veya teknik danışmanlık teşkil etmez; saha bazlı değerlendirme için yetkin bir danışmana başvurulmalıdır.
    </div>
  </footer>

  <script id="suayak-data" type="application/json">__SUAYAK_DATA_JSON__</script>
  <script>
  (function(){
    "use strict";
    var DATA = JSON.parse(document.getElementById('suayak-data').textContent);

    var SECTOR_COLORS = {
      "1. TEKSTİL VE HAZIR GİYİM": "#c084fc",
      "2. GIDA VE İÇECEK": "#eab308",
      "3. TARIM VE HAYVANCILIK": "#a3e635",
      "4. MADENCİLİK VE METALLER": "#a8a29e",
      "5. ENERJİ ÜRETİMİ": "#fb923c",
      "6. KAĞIT VE SELÜLOZ": "#4ade80",
      "7. KİMYA VE PETROKİMYA": "#f87171",
      "8. ELEKTRONİK VE YARI İLETKEN": "#818cf8",
      "9. DERİ İŞLEME (TABAKLAMA)": "#d97706",
      "10. VERİ MERKEZLERİ (SOĞUTMA)": "#38bdf8"
    };
    var INTENSITY_COLORS = {
      "Düşük": {bg:"rgba(74,222,128,.16)", fg:"#16a34a"},
      "Orta": {bg:"rgba(234,179,8,.16)", fg:"#a16207"},
      "Yüksek": {bg:"rgba(251,146,60,.18)", fg:"#c2410c"},
      "Çok Yüksek": {bg:"rgba(248,113,113,.18)", fg:"#b91c1c"}
    };
    var TYPE_COLORS = {
      "Mavi Su": {bg:"rgba(59,130,246,.16)", fg:"#1d4ed8"},
      "Yeşil Su": {bg:"rgba(74,222,128,.16)", fg:"#15803d"},
      "Gri Su": {bg:"rgba(148,163,184,.22)", fg:"#475569"},
      "Karma": {bg:"rgba(192,132,252,.18)", fg:"#7e22ce"}
    };

    var I18N = {
      tr: {
        title: "Su Ayak İzi Sektör Kataloğu",
        subtitle: "Sektöre göre su yoğunluğu, su ayak izi türü ve su riski; ISO 14046, AWS Standard, CDP Su Güvenliği ve GRI 303 çerçevelerine göre düzenlenmiştir",
        searchPlaceholder: "Sektör, süreç veya standart ara",
        allTypes: "Tüm Türler",
        excelBtn: "⬇ Excel İndir",
        recordWord: "kayıt",
        sectorWord: "sektör",
        sourceBadge: "✓ ISO 14046, AWS ve diğerleri",
        emptyMsg: "Aramanızla veya seçili sektörle eşleşen kayıt bulunamadı.",
        resetBtn: "Filtreleri temizle",
        allChip: "Tümü",
        standardsLabel: "Standartlar",
        mitigationLabel: "Azaltım",
        intensityLabel: "Yoğunluk"
      },
      en: {
        title: "Water Footprint Sector Catalogue",
        subtitle: "Water intensity, footprint type and water risk by sector, organised by ISO 14046, AWS Standard, CDP Water Security and GRI 303",
        searchPlaceholder: "Search sector, process or standard",
        allTypes: "All types",
        excelBtn: "⬇ Download Excel",
        recordWord: "records",
        sectorWord: "sectors",
        sourceBadge: "✓ ISO 14046, AWS and others",
        emptyMsg: "No records match your search or selected sector.",
        resetBtn: "Clear filters",
        allChip: "All",
        standardsLabel: "Standards",
        mitigationLabel: "Mitigation",
        intensityLabel: "Intensity"
      }
    };

    var sectors = [];
    DATA.forEach(function(d){ if(sectors.indexOf(d["Sektör"]) === -1) sectors.push(d["Sektör"]); });

    var urlParams = new URLSearchParams(window.location.search);
    var state = {
      sector: urlParams.get('sector') || "Tümü",
      type: urlParams.get('type') || "Tümü",
      query: urlParams.get('q') || ""
    };
    if(sectors.indexOf(state.sector) === -1) state.sector = "Tümü";

    var lang = "tr";
    try { lang = localStorage.getItem('suayak-lang') || "tr"; } catch(e){}

    var resultsEl = document.getElementById('results');
    var chipsEl = document.getElementById('sectorChips');
    var statsEl = document.getElementById('statsRow');
    var searchEl = document.getElementById('search');
    var clearBtn = document.getElementById('clearSearch');
    var typeFilterEl = document.getElementById('typeFilter');
    var exportBtn = document.getElementById('exportExcel');
    var themeBtn = document.getElementById('themeToggle');
    var langBtn = document.getElementById('langToggle');
    var titleEl = document.getElementById('brandTitle');
    var subtitleEl = document.getElementById('brandSubtitle');
    var typeAllOpt = document.getElementById('typeAllOpt');

    function shortLabel(s){ return s.replace(/^\d+\.\s*/, ''); }
    function esc(s){
      return String(s).replace(/[&<>]/g, function(c){
        return ({"&":"&amp;","<":"&lt;",">":"&gt;"})[c];
      });
    }
    function t(key){ return (I18N[lang] || I18N.tr)[key]; }

    function syncUrl(){
      var params = new URLSearchParams();
      if(state.sector !== 'Tümü') params.set('sector', state.sector);
      if(state.type !== 'Tümü') params.set('type', state.type);
      if(state.query) params.set('q', state.query);
      var qs = params.toString();
      var newUrl = window.location.pathname + (qs ? ('?' + qs) : '');
      window.history.replaceState(null, '', newUrl);
    }

    function renderChipsOnce(){
      var allActive = state.sector === 'Tümü';
      var html = '<button class="chip' + (allActive ? ' active' : '') + '" data-s="Tümü" aria-pressed="' + allActive + '"><span class="count">' + DATA.length + '</span> ' + esc(t('allChip')) + '</button>';
      sectors.forEach(function(s){
        var n = DATA.filter(function(d){ return d["Sektör"] === s; }).length;
        var color = SECTOR_COLORS[s] || '#94a3b8';
        var active = s === state.sector;
        html += '<button class="chip' + (active ? ' active' : '') + '" data-s="' + esc(s) + '" aria-pressed="' + active + '">' +
                '<span class="dot" style="background:' + color + '"></span>' +
                esc(shortLabel(s)) + '<span class="count">' + n + '</span></button>';
      });
      chipsEl.innerHTML = html;
      Array.prototype.forEach.call(chipsEl.querySelectorAll('.chip'), function(btn){
        btn.addEventListener('click', function(){
          state.sector = btn.getAttribute('data-s');
          Array.prototype.forEach.call(chipsEl.querySelectorAll('.chip'), function(b){
            var isActive = b === btn;
            b.classList.toggle('active', isActive);
            b.setAttribute('aria-pressed', isActive);
          });
          syncUrl();
          renderResults();
        });
      });
    }

    function matches(row){
      if(state.sector !== 'Tümü' && row["Sektör"] !== state.sector) return false;
      if(state.type !== 'Tümü' && row["Su Ayak İzi Türü"] !== state.type) return false;
      if(!state.query) return true;
      var q = state.query.toLocaleLowerCase('tr-TR');
      var hay = (
        row["Süreç"] + ' ' + row["İlgili Standartlar"] + ' ' + row["Başlıca Su Riski"] + ' ' +
        row["Azaltım Uygulamaları"] + ' ' + row["Açıklama"] + ' ' + row["Sektör"]
      ).toLocaleLowerCase('tr-TR');
      return hay.indexOf(q) !== -1;
    }

    function intensityBadge(level){
      var c = INTENSITY_COLORS[level] || {bg:"rgba(148,163,184,.16)", fg:"#64748b"};
      return '<span class="badge" style="background:' + c.bg + ';color:' + c.fg + '">' + esc(t('intensityLabel')) + ': ' + esc(level) + '</span>';
    }
    function typeBadge(type){
      var c = TYPE_COLORS[type] || {bg:"rgba(148,163,184,.16)", fg:"#64748b"};
      return '<span class="badge" style="background:' + c.bg + ';color:' + c.fg + '">' + esc(type) + '</span>';
    }

    function cardHTML(row){
      var color = SECTOR_COLORS[row["Sektör"]] || '#0891b2';
      return '<div class="card" style="--sc:' + color + '">' +
        '<div class="card-top">' +
          '<span class="card-cat"><span class="dot" style="background:' + color + '"></span>' + esc(shortLabel(row["Sektör"])) + '</span>' +
          typeBadge(row["Su Ayak İzi Türü"]) +
        '</div>' +
        '<div class="card-title">' + esc(row["Süreç"]) + '</div>' +
        '<div class="card-desc">' + esc(row["Açıklama"]) + '</div>' +
        '<div class="card-meta">' + intensityBadge(row["Su Yoğunluğu"]) + '</div>' +
        '<span class="detail"><b>' + esc(t('standardsLabel')) + ':</b> ' + esc(row["İlgili Standartlar"]) + '</span>' +
        '<span class="detail"><b>' + esc(t('mitigationLabel')) + ':</b> ' + esc(row["Azaltım Uygulamaları"]) + '</span>' +
      '</div>';
    }

    function renderResults(){
      var filtered = DATA.filter(matches);
      statsEl.innerHTML =
        '<div class="stat-pill"><b>' + filtered.length + '</b>&nbsp;' + esc(t('recordWord')) + '</div>' +
        '<div class="stat-pill"><b>' + sectors.length + '</b>&nbsp;' + esc(t('sectorWord')) + '</div>' +
        '<div class="stat-pill" style="color:#0e7490">' + esc(t('sourceBadge')) + '</div>';

      if(filtered.length === 0){
        resultsEl.innerHTML =
          '<div class="empty">' +
            '<svg width="34" height="34" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="1.5" stroke-linecap="round" stroke-linejoin="round"><circle cx="11" cy="11" r="8"/><line x1="21" y1="21" x2="16.65" y2="16.65"/></svg>' +
            '<p>' + esc(t('emptyMsg')) + '</p>' +
            '<button id="resetFilters">' + esc(t('resetBtn')) + '</button>' +
          '</div>';
        document.getElementById('resetFilters').addEventListener('click', function(){
          state.sector = 'Tümü'; state.query = ''; state.type = 'Tümü';
          searchEl.value = '';
          typeFilterEl.value = 'Tümü';
          Array.prototype.forEach.call(chipsEl.querySelectorAll('.chip'), function(b){
            var isAll = b.getAttribute('data-s') === 'Tümü';
            b.classList.toggle('active', isAll);
            b.setAttribute('aria-pressed', isAll);
          });
          clearBtn.classList.remove('show');
          syncUrl();
          renderResults();
        });
      } else {
        resultsEl.innerHTML = filtered.map(cardHTML).join('');
      }
      clearBtn.classList.toggle('show', !!state.query);
    }

    function applyTheme(theme){
      document.documentElement.setAttribute('data-theme', theme);
      themeBtn.textContent = theme === 'dark' ? '☀️' : '🌙';
    }

    function applyLang(){
      titleEl.textContent = t('title');
      subtitleEl.textContent = t('subtitle');
      searchEl.placeholder = t('searchPlaceholder');
      typeAllOpt.textContent = t('allTypes');
      exportBtn.textContent = t('excelBtn');
      langBtn.textContent = lang === 'tr' ? 'TR' : 'EN';
      document.documentElement.lang = lang;
      renderChipsOnce();
      renderResults();
    }

    var debounceTimer;
    searchEl.value = state.query;
    searchEl.addEventListener('input', function(){
      clearTimeout(debounceTimer);
      var val = searchEl.value;
      debounceTimer = setTimeout(function(){
        state.query = val;
        syncUrl();
        renderResults();
      }, 120);
    });

    clearBtn.addEventListener('click', function(){
      searchEl.value = '';
      state.query = '';
      syncUrl();
      renderResults();
      searchEl.focus();
    });

    typeFilterEl.value = state.type;
    typeFilterEl.addEventListener('change', function(){
      state.type = typeFilterEl.value;
      syncUrl();
      renderResults();
    });

    exportBtn.addEventListener('click', function(){
      var params = new URLSearchParams();
      if(state.sector !== 'Tümü') params.set('sector', state.sector);
      if(state.type !== 'Tümü') params.set('type', state.type);
      if(state.query) params.set('q', state.query);
      var qs = params.toString();
      window.location.href = '/export/xlsx' + (qs ? ('?' + qs) : '');
    });

    var savedTheme = null;
    try { savedTheme = localStorage.getItem('suayak-theme'); } catch(e){}
    applyTheme(savedTheme === 'dark' ? 'dark' : 'light');
    themeBtn.addEventListener('click', function(){
      var cur = document.documentElement.getAttribute('data-theme') === 'dark' ? 'light' : 'dark';
      applyTheme(cur);
      try { localStorage.setItem('suayak-theme', cur); } catch(e){}
    });

    langBtn.addEventListener('click', function(){
      lang = lang === 'tr' ? 'en' : 'tr';
      try { localStorage.setItem('suayak-lang', lang); } catch(e){}
      applyLang();
    });

    applyLang();

    if('serviceWorker' in navigator){
      window.addEventListener('load', function(){
        navigator.serviceWorker.register('/sw.js').catch(function(){});
      });
    }
  })();
  </script>
</body>
</html>
"""

_MANIFEST_JSON = r"""{
  "name": "Su Ayak İzi Sektör Kataloğu",
  "short_name": "Su Ayak İzi",
  "description": "Sektöre göre su yoğunluğu, su ayak izi türü ve su riski kataloğu",
  "start_url": "/",
  "scope": "/",
  "display": "standalone",
  "orientation": "portrait-primary",
  "background_color": "#eaf7fb",
  "theme_color": "#eaf7fb",
  "lang": "tr",
  "icons": [
    {"src": "/icons/icon-192.png", "sizes": "192x192", "type": "image/png", "purpose": "any"},
    {"src": "/icons/icon-512.png", "sizes": "512x512", "type": "image/png", "purpose": "any"},
    {"src": "/icons/icon-512-maskable.png", "sizes": "512x512", "type": "image/png", "purpose": "maskable"}
  ]
}
"""

_SERVICE_WORKER_JS = r"""// Su Ayak Izi Kataloglu servis calisani.
// Sayfa ve statik simgeler onbelleklenir; veri sayfaya gomulu oldugundan
// uygulama tamamen cevrimdisi calisabilir.
var CACHE_NAME = "suayak-cache-v1";
var APP_SHELL = ["/", "/manifest.json", "/icons/icon-192.png", "/icons/icon-512.png"];

self.addEventListener("install", function(e){
  e.waitUntil(
    caches.open(CACHE_NAME).then(function(cache){ return cache.addAll(APP_SHELL); })
  );
  self.skipWaiting();
});

self.addEventListener("activate", function(e){
  e.waitUntil(
    caches.keys().then(function(names){
      return Promise.all(names.filter(function(n){ return n !== CACHE_NAME; }).map(function(n){ return caches.delete(n); }));
    }).then(function(){ return self.clients.claim(); })
  );
});

self.addEventListener("fetch", function(e){
  if(e.request.method !== "GET"){ return; }
  e.respondWith(
    fetch(e.request).then(function(resp){
      var copy = resp.clone();
      caches.open(CACHE_NAME).then(function(cache){ cache.put(e.request, copy); });
      return resp;
    }).catch(function(){
      return caches.match(e.request).then(function(cached){ return cached || caches.match("/"); });
    })
  );
});
"""

_ERROR_404_HTML = r"""<!doctype html>
<html lang="tr" data-theme="light">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>404 - Sayfa Bulunamadı | Su Ayak İzi Kataloğu</title>
<style>
  body{margin:0;min-height:100vh;display:flex;align-items:center;justify-content:center;
    background:#eaf7fb;color:#0b2530;font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',Roboto,sans-serif;text-align:center;padding:24px;}
  .box{max-width:380px;}
  h1{font-size:56px;margin:0 0 8px;color:#0891b2;}
  p{font-size:14px;color:#35606b;line-height:1.6;margin:0 0 20px;}
  a{display:inline-block;background:#ffffff;border:1px solid #cfe8f0;color:#0b2530;
    padding:10px 18px;border-radius:8px;text-decoration:none;font-size:13px;}
</style>
</head>
<body>
  <div class="box">
    <h1>404</h1>
    <p>Aradığınız sayfa bulunamadı. Bağlantı hatalı olabilir veya sayfa taşınmış olabilir.</p>
    <a href="/">Ana sayfaya dön</a>
  </div>
</body>
</html>
"""

_ERROR_500_HTML = r"""<!doctype html>
<html lang="tr" data-theme="light">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>500 - Sunucu Hatası | Su Ayak İzi Kataloğu</title>
<style>
  body{margin:0;min-height:100vh;display:flex;align-items:center;justify-content:center;
    background:#eaf7fb;color:#0b2530;font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',Roboto,sans-serif;text-align:center;padding:24px;}
  .box{max-width:380px;}
  h1{font-size:56px;margin:0 0 8px;color:#d97706;}
  p{font-size:14px;color:#35606b;line-height:1.6;margin:0 0 20px;}
  a{display:inline-block;background:#ffffff;border:1px solid #cfe8f0;color:#0b2530;
    padding:10px 18px;border-radius:8px;text-decoration:none;font-size:13px;}
</style>
</head>
<body>
  <div class="box">
    <h1>500</h1>
    <p>Beklenmeyen bir sunucu hatası oluştu. Sorun devam ederse lütfen tekrar deneyin.</p>
    <a href="/">Ana sayfaya dön</a>
  </div>
</body>
</html>
"""


def _filter_data(sector=None, water_type=None, q=None):
    rows = DATA
    if sector and sector != "Tümü":
        rows = [r for r in rows if r["Sektör"] == sector]
    if water_type and water_type != "Tümü":
        rows = [r for r in rows if r["Su Ayak İzi Türü"] == water_type]
    if q:
        ql = q.lower()
        rows = [
            r for r in rows
            if ql in " ".join([
                r["Süreç"], r["İlgili Standartlar"], r["Başlıca Su Riski"],
                r["Azaltım Uygulamaları"], r["Açıklama"], r["Sektör"],
            ]).lower()
        ]
    return rows


@app.route("/")
def index():
    html = _PAGE_HTML.replace("__SUAYAK_DATA_JSON__", json.dumps(DATA, ensure_ascii=False))
    return Response(html, mimetype="text/html; charset=utf-8")


@app.route("/api/sectors")
def api_sectors():
    """Programatik erişim için JSON API.

    Opsiyonel sorgu parametreleri: ?sector=..., ?type=..., ?q=...
    """
    sector = request.args.get("sector")
    water_type = request.args.get("type")
    q = request.args.get("q")
    rows = _filter_data(sector=sector, water_type=water_type, q=q)
    return jsonify({
        "count": len(rows), "total": len(DATA),
        "sectors": SECTORS, "types": WATER_TYPES, "results": rows,
    })


@app.route("/export/xlsx")
def export_xlsx():
    sector = request.args.get("sector")
    water_type = request.args.get("type")
    q = request.args.get("q")
    rows = _filter_data(sector=sector, water_type=water_type, q=q)
    buf = _build_report_xlsx(rows, sector, water_type, q)
    return Response(
        buf.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="su-ayak-izi-katalog.xlsx"'},
    )


@app.route("/manifest.json")
def manifest():
    return Response(_MANIFEST_JSON, mimetype="application/manifest+json")


@app.route("/sw.js")
def service_worker():
    return Response(_SERVICE_WORKER_JS, mimetype="application/javascript")


@app.route("/icons/<path:filename>")
def icons(filename):
    if not os.path.isdir(ICONS_DIR):
        return Response(status=404)
    return send_from_directory(ICONS_DIR, filename)


@app.errorhandler(404)
def not_found(_e):
    return Response(_ERROR_404_HTML, status=404, mimetype="text/html; charset=utf-8")


@app.errorhandler(500)
def server_error(e):
    log.exception("Sunucu hatası: %s", e)
    return Response(_ERROR_500_HTML, status=500, mimetype="text/html; charset=utf-8")


# =============================================================================
# 3. "python suayakizi.py" ILE DOGRUDAN CALISTIRMA
# =============================================================================
def _open_browser():
    try:
        webbrowser.open(f"http://127.0.0.1:{PORT}")
    except Exception:
        pass


if __name__ == "__main__":
    print("=" * 58)
    print(" Su Ayak İzi Sektör Kataloğu")
    print(f" Sunucu adresi   : http://{HOST}:{PORT}")
    print(f" Kayıt sayısı    : {len(DATA)}")
    print(" Durdurmak için  : Ctrl+C")
    print("=" * 58)

    if HOST in ("127.0.0.1", "localhost"):
        Timer(0.6, _open_browser).start()

    try:
        app.run(host=HOST, port=PORT, debug=False)
    except OSError as exc:
        log.error("Sunucu başlatılamadı: %s", exc)
        log.error(
            "(%d numaralı port başka bir uygulama tarafından kullanılıyor olabilir; "
            "SUAYAK_PORT ortam değişkeniyle farklı bir port belirtebilirsiniz.)",
            PORT,
        )
